"""Build excel/nav_facility_stress.xlsx: the NAV facility model in formulas.

Run from the repo root:  python excel/build_nav_workbook.py
Needs openpyxl (not a package dependency). Excel recalculates on open.
Inputs are hypothetical and match examples/nav_facility_covenants_hypothetical.json.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
CFG = json.loads((ROOT / "examples" / "nav_facility_covenants_hypothetical.json").read_text(encoding="utf-8"))
OUT = ROOT / "excel" / "nav_facility_stress.xlsx"
N_ROWS, FIRST, LAST = 20, 17, 36            # asset rows 17..36
TOTAL, ELIG = 38, 42                        # summary rows

F = "Arial"
BLUE, BLACK, GREEN = Font(name=F, color="0000FF"), Font(name=F), Font(name=F, color="008000")
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="999999"))
PCT, NUM = "0.0%", "#,##0.0;(#,##0.0);-"

# Python model results for the shipped inputs (fund-finance-lab nav --config examples/nav_facility_covenants_hypothetical.json)
PY = {"eligible": 876.0, "ltv": 150 / 876, "to_sweep": 1 - 150 / (0.20 * 876), "to_breach": 1 - 150 / (0.25 * 876),
      "k1": 676.0, "k2": 516.0, "k3": 382.8, "status": "ok"}


def style(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font():                         # untouched cells get Arial too
                c.font = BLACK


def put(ws, ref, value, font=BLACK, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font = value, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


wb = Workbook()

# ---------------------------------------------------------------- README
rd = wb.active
rd.title = "README"
lines = [
    ("NAV facility stress model (formula version)", TITLE),
    ("Hypothetical inputs for illustration. Not observed market terms, not advice, and not any rating agency's methodology.", BOLD),
    ("", None),
    ("What it does", BOLD),
    ("Computes eligible NAV after a single-asset cap and a top-N concentration limit, LTV, the status against the cash-sweep", None),
    ("and breach levels, a diversity covenant, the uniform drawdown to each level, the cure to target, and the effect of", None),
    ("writing off the largest assets. It reproduces the Python model in src/fundfinancelab/nav_facility.py.", None),
    ("", None),
    ("How to use", BOLD),
    ("Edit only the blue cells on Inputs (facility terms and up to 20 assets). Yellow marks the key assumptions.", None),
    ("Leave Top-N count and share blank to switch that limit off; leave Minimum assets blank to switch the covenant off.", None),
    ("Every other cell is a formula. Checks shows whether the workbook is internally consistent.", None),
    ("", None),
    ("Sheets", BOLD),
    ("Inputs - facility terms and portfolio. Calc - eligible NAV, LTV, status, headroom. Uniform stress - drawdown grid.", None),
    ("Single-name - the largest 1, 2 and 3 assets written down by the loss on Inputs. Checks - consistency and Python cross-check.", None),
    ("", None),
    ("Method (see docs/model_spec.md)", BOLD),
    ("Single-asset cap: each asset counts up to cap x total NAV. Top-N limit: the N largest capped values count for at most", None),
    ("share x the capped total; the excess is excluded. Each limit is measured against the aggregate before it is applied.", None),
    ("A uniform drawdown d scales every NAV, cap and limit by (1 - d), so eligible NAV at d is eligible NAV x (1 - d).", None),
    ("Status: breach if LTV >= breach level; cash sweep if LTV >= sweep level or the diversity covenant fails; else ok.", None),
]
for i, (text, font) in enumerate(lines, start=1):
    put(rd, f"A{i}", text, font or BLACK)

# ---------------------------------------------------------------- Inputs
inp = wb.create_sheet("Inputs")
put(inp, "A1", "Inputs (edit blue cells only)", TITLE)
put(inp, "B3", "Facility terms", BOLD)
fac = CFG["facility"]
terms = [(5, "Loan outstanding (USD m)", fac["loan"], NUM, "Hypothetical"),
         (6, "Cash-sweep LTV", fac["ltv_sweep"], PCT, "Sweep applies at or above this LTV"),
         (7, "Breach LTV", fac["ltv_breach"], PCT, "Covenant breach at or above this LTV"),
         (8, "Cure target LTV", fac["ltv_target"], PCT, "LTV the borrower must restore after a breach"),
         (9, "Single-asset cap (share of total NAV)", fac["single_asset_cap"], PCT, "Each asset counts up to this share"),
         (10, "Top-N count", fac["top_n"], "0", "Blank = no top-N limit"),
         (11, "Top-N share of eligible NAV", fac["top_n_share"], PCT, "Blank = no top-N limit"),
         (12, "Minimum assets (diversity covenant)", fac["min_assets"], "0", "Blank = no covenant"),
         (13, "Single-name loss (share written off)", CFG["single_name_loss"], PCT, "Used on Single-name")]
for row, label, value, fmt, note in terms:
    put(inp, f"B{row}", label)
    put(inp, f"C{row}", value, BLUE, fmt, YELLOW if row in (5, 6, 7, 9, 10, 11, 12) else None)
    put(inp, f"D{row}", note)
put(inp, "B15", "Portfolio (up to 20 assets; USD m; hypothetical)", BOLD)
put(inp, "B16", "Asset", BOLD); put(inp, "C16", "NAV (USD m)", BOLD)
for i in range(N_ROWS):
    r = FIRST + i
    a = CFG["assets"][i] if i < len(CFG["assets"]) else None
    put(inp, f"B{r}", a["name"] if a else None, BLUE)
    put(inp, f"C{r}", a["nav"] if a else None, BLUE, NUM)
put(inp, "F5", "Legend", BOLD)
put(inp, "F6", "Blue text = input", BLUE)
put(inp, "F7", "Yellow fill = key assumption", BLACK, None, YELLOW)
put(inp, "F8", "Source: examples/nav_facility_covenants_hypothetical.json (hypothetical)")
style(inp, {"A": 3, "B": 40, "C": 14, "D": 44, "F": 34})

# ---------------------------------------------------------------- Calc
calc = wb.create_sheet("Calc")
put(calc, "A1", "Eligible NAV, LTV and status", TITLE)
for col, head in zip("BCDEFG", ["Asset", "NAV", "NAV rank", "Capped NAV", "Capped rank", "In top N"]):
    put(calc, f"{col}16", head, BOLD)
for r in range(FIRST, LAST + 1):
    put(calc, f"B{r}", f'=IF(Inputs!B{r}="","",Inputs!B{r})', GREEN)
    put(calc, f"C{r}", f"=IF(ISNUMBER(Inputs!C{r}),Inputs!C{r},0)", GREEN, NUM)
    put(calc, f"D{r}", f"=IF(C{r}>0,RANK(C{r},$C${FIRST}:$C${LAST},0)+COUNTIF(C$16:C{r-1},C{r}),999)", BLACK, "0")
    put(calc, f"E{r}", f"=MIN(C{r},Inputs!$C$9*$C${TOTAL})", BLACK, NUM)
    put(calc, f"F{r}", f"=RANK(E{r},$E${FIRST}:$E${LAST},0)+COUNTIF(E$16:E{r-1},E{r})", BLACK, "0")
    put(calc, f"G{r}", f"=IF(AND(ISNUMBER(Inputs!$C$10),F{r}<=Inputs!$C$10),1,0)", BLACK, "0")
summary = [
    (38, "Total NAV", f"=SUM(C{FIRST}:C{LAST})", NUM),
    (39, "Eligible NAV after single-asset cap", f"=SUM(E{FIRST}:E{LAST})", NUM),
    (40, "Top-N capped total", f"=SUMPRODUCT(E{FIRST}:E{LAST},G{FIRST}:G{LAST})", NUM),
    (41, "Excluded by top-N limit", "=IF(AND(ISNUMBER(Inputs!C10),ISNUMBER(Inputs!C11)),MAX(0,C40-Inputs!C11*C39),0)", NUM),
    (42, "Eligible NAV", "=C39-C41", NUM),
    (43, "Excluded by single-asset cap", "=C38-C39", NUM),
    (44, "LTV", '=IF(C42>0,Inputs!C5/C42,"n/a")', PCT),
    (45, "Assets with value", f'=COUNTIF(C{FIRST}:C{LAST},">0")', "0"),
    (46, "Diversity covenant met", "=IF(ISNUMBER(Inputs!C12),C45>=Inputs!C12,TRUE)", None),
    (47, "Status", '=IF(C44>=Inputs!C7,"breach",IF(OR(C44>=Inputs!C6,NOT(C46)),"cash sweep","ok"))', None),
    (48, "Uniform drawdown to cash-sweep level", "=MAX(0,1-Inputs!C5/(Inputs!C6*C42))", PCT),
    (49, "Uniform drawdown to breach level", "=MAX(0,1-Inputs!C5/(Inputs!C7*C42))", PCT),
    (50, "Cure to target (if in breach)", "=IF(C44>=Inputs!C7,Inputs!C5-Inputs!C8*C42,0)", NUM),
]
for row, label, formula, fmt in summary:
    put(calc, f"B{row}", label, BOLD if row in (42, 44, 47) else BLACK)
    put(calc, f"C{row}", formula, BLACK, fmt)
put(calc, "B52", "NAV rank 999 marks an empty row. Ranks break ties by row order, as the Python model does.")
style(calc, {"A": 3, "B": 38, "C": 14, "D": 10, "E": 13, "F": 12, "G": 10})

# ---------------------------------------------------------------- Uniform stress
uni = wb.create_sheet("Uniform stress")
put(uni, "A1", "Uniform drawdowns", TITLE)
put(uni, "A3", "Eligible NAV at drawdown d = eligible NAV x (1 - d): every NAV, cap and limit scales together.")
for col, head in zip("ABCDE", ["Drawdown", "Eligible NAV", "LTV", "Status", "Cure to target"]):
    put(uni, f"{col}5", head, BOLD)
for i, d in enumerate(CFG["drawdowns"]):
    r = 6 + i
    put(uni, f"A{r}", d, BLUE, PCT)
    put(uni, f"B{r}", f"=Calc!$C$42*(1-A{r})", GREEN, NUM)
    put(uni, f"C{r}", f'=IF(B{r}>0,Inputs!$C$5/B{r},"n/a")', BLACK, PCT)
    put(uni, f"D{r}", f'=IF(C{r}>=Inputs!$C$7,"breach",IF(OR(C{r}>=Inputs!$C$6,NOT(Calc!$C$46)),"cash sweep","ok"))')
    put(uni, f"E{r}", f"=IF(C{r}>=Inputs!$C$7,Inputs!$C$5-Inputs!$C$8*B{r},0)", BLACK, NUM)
style(uni, {"A": 12, "B": 14, "C": 10, "D": 14, "E": 16})

# ---------------------------------------------------------------- Single-name
sn = wb.create_sheet("Single-name")
put(sn, "A1", "Largest assets written down", TITLE)
put(sn, "A3", "Block k writes down the k largest assets (by NAV) by the loss on Inputs; limits are recomputed on the smaller portfolio.")
put(sn, "B16", "Asset", BOLD)
blocks = {1: "CDEF", 2: "HIJK", 3: "MNOP"}
for k, (cs, cc, cr, cf) in blocks.items():
    put(sn, f"{cs}15", f"k = {k}", BOLD)
    put(sn, f"{cc}15", k, BLACK, "0")
    for col, head in zip((cs, cc, cr, cf), ("Stressed NAV", "Capped", "Rank", "Top N")):
        put(sn, f"{col}16", head, BOLD)
    for r in range(FIRST, LAST + 1):
        if k == 1:
            put(sn, f"B{r}", f"=Calc!B{r}", GREEN)
        put(sn, f"{cs}{r}", f"=IF(Calc!$D{r}<=${cc}$15,Calc!$C{r}*(1-Inputs!$C$13),Calc!$C{r})", BLACK, NUM)
        put(sn, f"{cc}{r}", f"=MIN({cs}{r},Inputs!$C$9*{cs}${TOTAL})", BLACK, NUM)
        put(sn, f"{cr}{r}", f"=RANK({cc}{r},{cc}${FIRST}:{cc}${LAST},0)+COUNTIF({cc}$16:{cc}{r-1},{cc}{r})", BLACK, "0")
        put(sn, f"{cf}{r}", f"=IF(AND(ISNUMBER(Inputs!$C$10),{cr}{r}<=Inputs!$C$10),1,0)", BLACK, "0")
    rows = [
        (38, "Total NAV", f"=SUM({cs}{FIRST}:{cs}{LAST})", NUM),
        (39, "Eligible after single-asset cap", f"=SUM({cc}{FIRST}:{cc}{LAST})", NUM),
        (40, "Top-N capped total", f"=SUMPRODUCT({cc}{FIRST}:{cc}{LAST},{cf}{FIRST}:{cf}{LAST})", NUM),
        (41, "Excluded by top-N limit", f"=IF(AND(ISNUMBER(Inputs!$C$10),ISNUMBER(Inputs!$C$11)),MAX(0,{cs}40-Inputs!$C$11*{cs}39),0)", NUM),
        (42, "Eligible NAV", f"={cs}39-{cs}41", NUM),
        (43, "LTV", f'=IF({cs}42>0,Inputs!$C$5/{cs}42,"n/a")', PCT),
        (44, "Assets with value", f'=COUNTIF({cs}{FIRST}:{cs}{LAST},">0")', "0"),
        (45, "Diversity covenant met", f"=IF(ISNUMBER(Inputs!$C$12),{cs}44>=Inputs!$C$12,TRUE)", None),
        (46, "Status", f'=IF({cs}43>=Inputs!$C$7,"breach",IF(OR({cs}43>=Inputs!$C$6,NOT({cs}45)),"cash sweep","ok"))', None),
        (47, "Further uniform drawdown to breach", f"=MAX(0,1-Inputs!$C$5/(Inputs!$C$7*{cs}42))", PCT),
    ]
    for row, label, formula, fmt in rows:
        if k == 1:
            put(sn, f"B{row}", label, BOLD if row in (42, 43, 46) else BLACK)
        put(sn, f"{cs}{row}", formula, BLACK, fmt)
style(sn, {"A": 3, "B": 34, **{c: 12 for c in "CDEFHIJKMNOP"}, "G": 3, "L": 3})

# ---------------------------------------------------------------- Checks
ck = wb.create_sheet("Checks")
put(ck, "A1", "Checks", TITLE)
put(ck, "A3", "Internal consistency (must all be TRUE)", BOLD)
checks = [
    ("Cure target <= sweep level <= breach level < 100%", "=AND(Inputs!C8<=Inputs!C6,Inputs!C6<=Inputs!C7,Inputs!C7<1)"),
    ("No negative NAVs", f'=COUNTIF(Inputs!C{FIRST}:C{LAST},"<0")=0'),
    ("Eligible NAV <= total NAV", "=Calc!C42<=Calc!C38"),
    ("Capped ranks are 1..20 with no ties", f"=AND(SUM(Calc!F{FIRST}:F{LAST})=210,MIN(Calc!F{FIRST}:F{LAST})=1,MAX(Calc!F{FIRST}:F{LAST})=20)"),
    # Drawdown to breach is floored at 0 (already in breach) and is 100% only when the loan is 0; the identity is
    # checked only in between. Found by the regression cases (excel/regression): the unguarded version failed for
    # every portfolio already in breach and returned #DIV/0! for a zero loan.
    ("LTV at the breakeven drawdown equals the breach level (or LTV is already at or above it)",
     "=IF(Calc!C49>=1,Inputs!C5=0,IF(Calc!C49>0,ABS(Inputs!C5/(Calc!C42*(1-Calc!C49))-Inputs!C7)<0.000000001,"
     "Calc!C44>=Inputs!C7))"),
    ("Top-N exclusion is not negative", "=Calc!C41>=0"),
]
for i, (label, formula) in enumerate(checks):
    put(ck, f"A{5 + i}", label)
    put(ck, f"B{5 + i}", formula)
put(ck, "A12", "Overall", BOLD)
put(ck, "B12", '=IF(AND(B5:B10),"ALL CHECKS PASS","CHECK FAILED")', BOLD)
put(ck, "A15", "Cross-check against the Python model (shipped inputs only)", BOLD)
put(ck, "A16", "Reference values from: fund-finance-lab nav --config examples/nav_facility_covenants_hypothetical.json. "
               "They are expected to differ once inputs change.")
for col, head in zip("ABCD", ["Metric", "Python", "Workbook", "Match"]):
    put(ck, f"{col}18", head, BOLD)
xref = [("Eligible NAV", PY["eligible"], "=Calc!C42", NUM), ("LTV", PY["ltv"], "=Calc!C44", "0.0000%"),
        ("Uniform drawdown to sweep", PY["to_sweep"], "=Calc!C48", "0.0000%"),
        ("Uniform drawdown to breach", PY["to_breach"], "=Calc!C49", "0.0000%"),
        ("Eligible NAV, largest asset written off", PY["k1"], "='Single-name'!C42", NUM),
        ("Eligible NAV, two largest written off", PY["k2"], "='Single-name'!H42", NUM),
        ("Eligible NAV, three largest written off", PY["k3"], "='Single-name'!M42", NUM)]
for i, (label, py, formula, fmt) in enumerate(xref):
    r = 19 + i
    put(ck, f"A{r}", label)
    put(ck, f"B{r}", py, BLUE, fmt)
    put(ck, f"C{r}", formula, GREEN, fmt)
    put(ck, f"D{r}", f'=IF(ABS(C{r}-B{r})<0.000001,"match","MISMATCH")')
r = 19 + len(xref)
put(ck, f"A{r}", "Status at base")
put(ck, f"B{r}", PY["status"], BLUE)
put(ck, f"C{r}", "=Calc!C47", GREEN)
put(ck, f"D{r}", f'=IF(C{r}=B{r},"match","MISMATCH")')
style(ck, {"A": 52, "B": 16, "C": 16, "D": 12})

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title != "README"
    for row in ws.iter_rows():
        for c in row:
            if c.alignment == Alignment():
                c.alignment = Alignment(vertical="center")
wb.save(OUT)
print(f"wrote {OUT}")
