"""Evaluate excel/nav_facility_stress.xlsx on every regression case with LibreOffice.

Run from the repo root:  python excel/regression/run_libreoffice.py
Needs openpyxl and LibreOffice (soffice on PATH). Not part of CI.

For each row of excel/regression/cases.csv it writes the inputs into a copy of the workbook,
recalculates all copies in one LibreOffice session, reads the output cells, and writes
excel/regression/libreoffice_log.csv in the same format the VBA harness writes. The log is then
checked with:  fund-finance-lab verify-log --log excel/regression/libreoffice_log.csv

This is LibreOffice's evaluation of the formulas, not Excel's. The Excel run is the VBA harness.
"""
from __future__ import annotations

import csv
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from fundfinancelab import regression as reg  # noqa: E402

WORKBOOK = ROOT / "excel" / "nav_facility_stress.xlsx"
CASES = ROOT / "excel" / "regression" / "cases.csv"
LOG = ROOT / "excel" / "regression" / "libreoffice_log.csv"

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalcOne(sURL As String)
    Dim oDoc As Object
    Dim args(0) As New com.sun.star.beans.PropertyValue
    args(0).Name = "Hidden"
    args(0).Value = True
    oDoc = StarDesktop.loadComponentFromURL(sURL, "_blank", 0, args())
    oDoc.calculateAll()
    oDoc.store()
    oDoc.close(True)
End Sub
Sub RecalcAll
{calls}
    StarDesktop.terminate()
End Sub
</script:module>
"""


def split(ref: str) -> tuple[str, str]:
    sheet, cell = ref.rsplit("!", 1)
    return sheet.strip("'"), cell


def recalc_all(files: list[Path], timeout: int = 900) -> None:
    """Recalculate and re-save every workbook in one LibreOffice session."""
    env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        url = Path(profile).as_uri()
        subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       env=env, capture_output=True, timeout=180, check=False)
        macro_dir = Path(profile) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            raise RuntimeError("LibreOffice did not create a profile; nothing was recalculated")
        calls = "\n".join(f'    RecalcOne("{f.resolve().as_uri()}")' for f in files)
        (macro_dir / "Module1.xba").write_text(MACRO.format(calls=calls), encoding="utf-8")
        before = {f: f.stat().st_mtime_ns for f in files}
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
                        "vnd.sun.star.script:Standard.Module1.RecalcAll?language=Basic&location=application"],
                       env=env, capture_output=True, timeout=timeout, check=False)
        stale = [f.name for f in files if f.stat().st_mtime_ns == before[f]]
        if stale:
            raise RuntimeError(f"LibreOffice did not rewrite {len(stale)} file(s), e.g. {stale[:3]}")


def as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return repr(float(v))
    return str(v)


def main() -> int:
    rows = reg.read_cases(CASES)
    with tempfile.TemporaryDirectory(prefix="nav-cases-") as tmp:
        files = []
        for row in rows:
            case = reg.case_inputs(row)
            wb = load_workbook(WORKBOOK)
            for field, ref in reg.INPUT_CELLS.items():
                sheet, cell = split(ref)
                wb[sheet][cell] = case[field]
            for i, (nref, vref) in enumerate(zip(reg.NAME_CELLS, reg.NAV_CELLS), 1):
                wb["Inputs"][split(nref)[1]] = case[f"name_{i:02d}"]
                wb["Inputs"][split(vref)[1]] = case[f"nav_{i:02d}"]
            path = Path(tmp) / f"{row['case_id']}.xlsx"
            wb.save(path)
            files.append(path)
        recalc_all(files)
        with open(LOG, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(reg.LOG_COLUMNS)
            for row, path in zip(rows, files):
                wb = load_workbook(path, data_only=True)
                for field, ref, _ in reg.OUTPUTS:
                    sheet, cell = split(ref)
                    actual = as_text(wb[sheet][cell].value)
                    ok = reg.matches(field, row[f"exp_{field}"], actual)
                    w.writerow([row["case_id"], field, actual, "TRUE" if ok else "FALSE"])
    result = reg.verify_log(CASES, LOG)
    print(reg.report_text(result, LOG.name), end="")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
